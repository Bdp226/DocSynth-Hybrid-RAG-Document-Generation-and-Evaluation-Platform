from __future__ import annotations

import base64
import binascii
from io import BytesIO
from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    Image as RLImage,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .models import ImageInput


@dataclass
class RenderedArtifact:
    format: str
    filename: str
    mime_type: str
    content: bytes


# Brand palette matches the UI theme (black/navy/teal/orange).
BRAND_NAVY = "#10213A"
BRAND_TEAL = "#009F93"
BRAND_ORANGE = "#FF7A1A"
LOGO_CANDIDATE_PATHS = (
    Path(__file__).parent / "static" / "branding" / "siemens_healthineers_logo.png",
    Path(__file__).parent / "static" / "branding" / "siemens_healthineers_logo.jpg",
    Path(__file__).parent / "static" / "branding" / "siemens_healthineers_logo.jpeg",
)

_FULL_BOLD_LINE_RE = re.compile(r"^\*\*(.+)\*\*$")
_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_INLINE_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)([^*]+?)(?<!\*)\*(?!\*)")
_IMAGE_TOKEN_RE = re.compile(r"^\[\[IMAGE:(.+?)\]\]$")


def _figure_caption(index: int, section_title: str = "") -> str:
    """Build a textbook-style caption that never exposes source slide numbering."""
    if section_title:
        return f"Figure {index} — {section_title}"
    return f"Figure {index}"


def _image_lookup(image_inputs: list[ImageInput] | None) -> dict[str, bytes]:
    lookup: dict[str, bytes] = {}
    for image in image_inputs or []:
        content = _normalize_base64_image(image.content_base64)
        if content is not None:
            lookup[image.image_name] = content
    return lookup


def _is_table_row(line: str) -> bool:
    return line.startswith("|") and line.endswith("|") and line.count("|") >= 3


def _is_table_divider(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", cell.strip()) for cell in cells)


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _normalize_table_rows(rows: list[str]) -> list[list[str]]:
    """Convert buffered markdown rows into a rectangular matrix, dropping divider rows."""
    matrix = [_split_table_row(row) for row in rows if not _is_table_divider(row)]
    matrix = [row for row in matrix if any(cell for cell in row)]
    if not matrix:
        return []
    width = max(len(row) for row in matrix)
    return [row + [""] * (width - len(row)) for row in matrix]


def _escape_xml(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _markdown_to_reportlab_markup(value: str) -> str:
    escaped = _escape_xml(value)
    escaped = _INLINE_BOLD_RE.sub(r"<b>\1</b>", escaped)
    escaped = _INLINE_ITALIC_RE.sub(r"<i>\1</i>", escaped)
    return escaped


def _resolve_logo_path() -> Path | None:
    for path in LOGO_CANDIDATE_PATHS:
        if path.exists() and path.is_file():
            return path
    return None


def _normalize_base64_image(content: str) -> bytes | None:
    marker = "base64,"
    normalized = content.split(marker, 1)[1].strip() if marker in content else content.strip()
    try:
        return base64.b64decode(normalized, validate=True)
    except (binascii.Error, ValueError):
        return None


def _draw_brand_header(canvas_obj, doc_obj) -> None:
    page_width, page_height = A4
    right_x = page_width - 40
    top_y = page_height - 34
    logo_path = _resolve_logo_path()

    canvas_obj.saveState()
    if logo_path is not None:
        logo = ImageReader(str(logo_path))
        logo_width, logo_height = logo.getSize()
        target_width = 128
        scale = target_width / float(logo_width)
        target_height = logo_height * scale
        canvas_obj.drawImage(
            logo,
            right_x - target_width,
            top_y - target_height + 6,
            width=target_width,
            height=target_height,
            mask="auto",
            preserveAspectRatio=True,
            anchor="ne",
        )
    else:
        canvas_obj.setFont("Helvetica-Bold", 13)
        canvas_obj.setFillColor(colors.HexColor(BRAND_NAVY))
        canvas_obj.drawRightString(right_x, top_y, "SIEMENS")
        canvas_obj.setFont("Helvetica-Bold", 13)
        canvas_obj.setFillColor(colors.HexColor(BRAND_ORANGE))
        canvas_obj.drawRightString(right_x, top_y - 15, "Healthineers")
    canvas_obj.setStrokeColor(colors.HexColor(BRAND_TEAL))
    canvas_obj.setLineWidth(1.1)
    canvas_obj.line(40, top_y - 24, page_width - 40, top_y - 24)
    canvas_obj.restoreState()


def _build_pdf_bytes(text: str, image_inputs: list[ImageInput] | None = None) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=48,
        rightMargin=48,
        topMargin=78,
        bottomMargin=48,
        title="Generated Document",
    )

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor(BRAND_NAVY),
        spaceAfter=8,
    )
    h1_style = ParagraphStyle(
        "H1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=colors.HexColor(BRAND_NAVY),
        spaceAfter=12,
        spaceBefore=4,
    )
    h2_style = ParagraphStyle(
        "H2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        textColor=colors.HexColor(BRAND_TEAL),
        spaceAfter=8,
        spaceBefore=10,
    )
    h3_style = ParagraphStyle(
        "H3",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        textColor=colors.HexColor(BRAND_NAVY),
        spaceAfter=6,
        spaceBefore=8,
    )

    caption_style = ParagraphStyle(
        "Caption",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor(BRAND_TEAL),
        spaceAfter=10,
    )

    story = []
    bullet_buffer: list[str] = []
    table_buffer: list[str] = []
    title_used = False
    figure_number = 0
    current_section = ""
    image_lookup = _image_lookup(image_inputs)
    used_image_names: set[str] = set()

    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1B1B1B"),
    )
    header_cell_style = ParagraphStyle(
        "TableHeaderCell",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    def flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        matrix = _normalize_table_rows(table_buffer)
        table_buffer = []
        if not matrix:
            return

        data = [
            [Paragraph(_markdown_to_reportlab_markup(cell), header_cell_style if row_index == 0 else cell_style) for cell in row]
            for row_index, row in enumerate(matrix)
        ]
        available_width = A4[0] - 96
        column_width = available_width / len(matrix[0])
        table = Table(data, colWidths=[column_width] * len(matrix[0]), repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_NAVY)),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F8")]),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B9C4D0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(Spacer(1, 6))
        story.append(table)
        story.append(Spacer(1, 10))

    def flush_bullets() -> None:
        nonlocal bullet_buffer
        if not bullet_buffer:
            return
        items = [ListItem(Paragraph(_markdown_to_reportlab_markup(item), body_style)) for item in bullet_buffer]
        story.append(ListFlowable(items, bulletType="bullet", start="circle", leftIndent=18))
        story.append(Spacer(1, 6))
        bullet_buffer = []

    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            flush_bullets()
            flush_table()
            story.append(Spacer(1, 6))
            continue

        if _is_table_row(line):
            flush_bullets()
            table_buffer.append(line)
            continue
        flush_table()

        image_token = _IMAGE_TOKEN_RE.match(line)
        if image_token:
            flush_bullets()
            image_name = image_token.group(1).strip()
            image_bytes = image_lookup.get(image_name)
            if image_bytes is not None:
                used_image_names.add(image_name)
                figure_number += 1
                story.append(Spacer(1, 6))
                story.append(RLImage(BytesIO(image_bytes), width=430, height=260, kind="proportional"))
                story.append(
                    Paragraph(
                        _markdown_to_reportlab_markup(_figure_caption(figure_number, current_section)),
                        caption_style,
                    )
                )
            continue

        if re.match(r"^[-*•]\s+", line):
            bullet_buffer.append(re.sub(r"^[-*•]\s+", "", line))
            continue

        flush_bullets()

        full_bold_match = _FULL_BOLD_LINE_RE.match(line)
        if line.startswith("# "):
            story.append(Paragraph(_markdown_to_reportlab_markup(line[2:].strip()), h1_style))
            title_used = True
        elif full_bold_match:
            heading_text = _markdown_to_reportlab_markup(full_bold_match.group(1).strip())
            if not title_used:
                story.append(Paragraph(heading_text, h1_style))
                title_used = True
            else:
                story.append(Paragraph(heading_text, h2_style))
        elif line.startswith("### "):
            heading = line[4:].strip()
            current_section = heading
            story.append(Paragraph(_markdown_to_reportlab_markup(heading), h3_style))
        elif line.startswith("## ") or line.endswith(":"):
            heading = line[3:].strip() if line.startswith("## ") else line
            current_section = heading
            story.append(Paragraph(_markdown_to_reportlab_markup(heading), h2_style))
        else:
            story.append(Paragraph(_markdown_to_reportlab_markup(line), body_style))

    flush_bullets()
    flush_table()
    remaining_images = [image for image in (image_inputs or []) if image.image_name not in used_image_names]
    if remaining_images:
        story.append(Spacer(1, 14))
        story.append(Paragraph("Image Appendix", h2_style))
        for image in remaining_images:
            image_bytes = image_lookup.get(image.image_name)
            if image_bytes is None:
                continue
            story.append(Paragraph(_markdown_to_reportlab_markup(image.image_name), body_style))
            story.append(Spacer(1, 4))
            story.append(RLImage(BytesIO(image_bytes), width=420, height=220, kind="proportional"))
            story.append(Spacer(1, 10))
    doc.build(story, onFirstPage=_draw_brand_header, onLaterPages=_draw_brand_header)
    return buffer.getvalue()


def _add_brand_header_docx(doc: Document) -> None:
    header = doc.sections[0].header
    header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_para.text = ""

    logo_path = _resolve_logo_path()
    if logo_path is not None:
        run = header_para.add_run()
        run.add_picture(str(logo_path), width=Inches(1.7))
        return

    siemens_run = header_para.add_run("SIEMENS ")
    siemens_run.bold = True
    siemens_run.font.size = Pt(14)
    siemens_run.font.color.rgb = RGBColor(0x10, 0x21, 0x3A)

    healthineers_run = header_para.add_run("Healthineers")
    healthineers_run.bold = True
    healthineers_run.font.size = Pt(14)
    healthineers_run.font.color.rgb = RGBColor(0xFF, 0x7A, 0x1A)


def _add_markdown_runs_docx(paragraph, text: str) -> None:
    tokens = re.split(r"(\*\*[^*]+?\*\*)", text)
    for token in tokens:
        if not token:
            continue
        bold_match = re.match(r"^\*\*(.+)\*\*$", token)
        run = paragraph.add_run(bold_match.group(1) if bold_match else token)
        run.bold = bool(bold_match)


def _build_docx_bytes(text: str, image_inputs: list[ImageInput] | None = None) -> bytes:
    doc = Document()
    _add_brand_header_docx(doc)

    bullet_buffer: list[str] = []
    table_buffer: list[str] = []
    title_used = False
    figure_number = 0
    current_section = ""
    image_lookup = _image_lookup(image_inputs)
    used_image_names: set[str] = set()

    def flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        matrix = _normalize_table_rows(table_buffer)
        table_buffer = []
        if not matrix:
            return

        table = doc.add_table(rows=len(matrix), cols=len(matrix[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(matrix):
            for column_index, cell_text in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.text = ""
                paragraph = cell.paragraphs[0]
                _add_markdown_runs_docx(paragraph, cell_text)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if row_index == 0:
                        run.bold = True
        doc.add_paragraph()

    def flush_bullets() -> None:
        nonlocal bullet_buffer
        for item in bullet_buffer:
            bullet_para = doc.add_paragraph(style="List Bullet")
            _add_markdown_runs_docx(bullet_para, item)
        bullet_buffer = []

    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            flush_bullets()
            flush_table()
            continue

        if _is_table_row(line):
            flush_bullets()
            table_buffer.append(line)
            continue
        flush_table()

        image_token = _IMAGE_TOKEN_RE.match(line)
        if image_token:
            flush_bullets()
            image_name = image_token.group(1).strip()
            image_bytes = image_lookup.get(image_name)
            if image_bytes is not None:
                used_image_names.add(image_name)
                figure_number += 1
                doc.add_picture(BytesIO(image_bytes), width=Inches(5.8))
                caption_para = doc.add_paragraph()
                caption_run = caption_para.add_run(_figure_caption(figure_number, current_section))
                caption_run.italic = True
                caption_run.font.size = Pt(9)
            continue

        if re.match(r"^[-*•]\s+", line):
            bullet_buffer.append(re.sub(r"^[-*•]\s+", "", line))
            continue

        flush_bullets()

        full_bold_match = _FULL_BOLD_LINE_RE.match(line)
        if line.startswith("# "):
            heading_para = doc.add_heading(level=1)
            _add_markdown_runs_docx(heading_para, line[2:].strip())
            title_used = True
        elif full_bold_match:
            heading_text = full_bold_match.group(1).strip()
            heading_para = doc.add_heading(level=1 if not title_used else 2)
            heading_para.add_run(heading_text).bold = True
            title_used = True
        elif line.startswith("### "):
            heading = line[4:].strip()
            current_section = heading
            heading_para = doc.add_heading(level=3)
            _add_markdown_runs_docx(heading_para, heading)
        elif line.startswith("## ") or line.endswith(":"):
            heading = line[3:].strip() if line.startswith("## ") else line
            current_section = heading
            heading_para = doc.add_heading(level=2)
            _add_markdown_runs_docx(heading_para, heading)
        else:
            body_para = doc.add_paragraph()
            _add_markdown_runs_docx(body_para, line)

    flush_bullets()
    flush_table()

    remaining_images = [image for image in (image_inputs or []) if image.image_name not in used_image_names]
    if remaining_images:
        doc.add_heading("Image Appendix", level=2)
        for image in remaining_images:
            image_bytes = image_lookup.get(image.image_name)
            if image_bytes is None:
                continue
            doc.add_paragraph(image.image_name)
            doc.add_picture(BytesIO(image_bytes), width=Inches(5.8))

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_artifacts(document_id: str, text: str, formats: list[str], image_inputs: list[ImageInput] | None = None) -> list[RenderedArtifact]:
    artifacts: list[RenderedArtifact] = []
    normalized = [fmt.strip().lower() for fmt in formats]

    if "pdf" in normalized:
        raw_pdf = _build_pdf_bytes(text, image_inputs=image_inputs)
        artifacts.append(
            RenderedArtifact(
                format="pdf",
                filename=f"{document_id}.pdf",
                mime_type="application/pdf",
                content=raw_pdf,
            )
        )

    if "docx" in normalized:
        raw_docx = _build_docx_bytes(text, image_inputs=image_inputs)
        artifacts.append(
            RenderedArtifact(
                format="docx",
                filename=f"{document_id}.docx",
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                content=raw_docx,
            )
        )

    return artifacts

from __future__ import annotations

import base64
import hashlib
import math
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import ImageInput


SUPPORTED_SUFFIXES = {".md", ".txt", ".docx", ".pdf", ".pptx"}

_EMBEDDING_MODEL = None
_WORKSPACE_TEXT_CACHE: dict[tuple[str, int], tuple[tuple[int, int, int], str]] = {}
_NOISE_LINE_PATTERNS = [
    re.compile(r"^author\s*\|\s*department$", re.IGNORECASE),
    re.compile(r"^shift\s+ise\s+india$", re.IGNORECASE),
    re.compile(r"^icb\s+and\s+tac\s+india$", re.IGNORECASE),
    re.compile(r"^icb\s+and\s+sc\s+india$", re.IGNORECASE),
    re.compile(r"^\[slide\s+\d+\]$", re.IGNORECASE),
    re.compile(r"^page\s+\d+$", re.IGNORECASE),
    re.compile(r"^topic\s*:\s*.*$", re.IGNORECASE),
    re.compile(r"^vendor\s+team\s+members$", re.IGNORECASE),
    re.compile(r"^shs\s+team\s+members$", re.IGNORECASE),
    re.compile(r"^last\s+updated\s*:\s*.*$", re.IGNORECASE),
    re.compile(r"^what\s+we\s+plan\s+to\s+do\??$", re.IGNORECASE),
    re.compile(r"^speaker\s+notes\s*:?$", re.IGNORECASE),
    re.compile(r"^(project|shift|vendor|shs)\s+team\s*:?.*$", re.IGNORECASE),
    re.compile(r"^project\s+owner\s+from\s+business.*$", re.IGNORECASE),
]

# Roles observed in the source decks, used to drop personnel rosters from documentation.
_ROLE_PAREN_RE = re.compile(
    r"\((?:founder|eng\.?\s*head|dev|po|ar|pl(?:\s*&\s*sm)?|sm|qa|lead|manager|director|architect|owner|cto|ceo)\)",
    re.IGNORECASE,
)
_PERSON_NAME_RE = re.compile(r"\b[A-Z][a-z]{2,}\s+[A-Z][a-zA-Z]{1,}\b")
# Title-case domain phrases look like names; these words mark a match as non-personal.
_NON_NAME_WORDS = {
    "access", "agent", "agents", "api", "architecture", "atlas", "budget", "cloud", "code", "compliance",
    "cost", "crew", "dashboard", "dashboards", "data", "design", "effort", "environment", "estimation",
    "evaluation", "flow", "gateway", "governance", "hub", "infrastructure", "machine", "management",
    "meeting", "model", "network", "onboarding", "overview", "phase", "pilot", "plan", "platform",
    "process", "project", "release", "report", "review", "risk", "sandbox", "security", "server",
    "service", "services", "setup", "solution", "status", "summary", "system", "team", "test",
    "testing", "tool", "tools", "update", "vendor", "version", "virtual",
}
# Chatter is a named participant plus a reporting verb, not the verb on its own.
_MEETING_CHATTER_RE = re.compile(
    r"(?i:\blooping\s+in\b)"
    r"|\b[A-Z][a-z]+(?:\s+[A-Z][a-z]+)?\s+(?:has\s+|have\s+|had\s+)?"
    r"(?:discussed|introduced|asked|explained|mentioned)\b",
)


@dataclass
class RetrievedChunk:
    source: str
    chunk_id: str
    score: float
    text: str


@dataclass
class RetrievalStats:
    candidate_files: int
    selected_files: int
    chunks_scored: int
    returned_chunks: int
    cache_hits: int
    cache_misses: int
    context_chars: int
    retrieval_latency_ms: int
    top_score: float
    avg_top_score: float
    used_embeddings: bool


def _tokenize(text: str) -> set[str]:
    return {t for t in re.findall(r"[a-zA-Z0-9_]+", text.lower()) if len(t) > 2}


def _personal_name_matches(line: str) -> list[str]:
    """Return name-like matches, excluding title-case domain phrases."""
    matches = []
    for match in _PERSON_NAME_RE.findall(line):
        tokens = {token.lower().strip(".,;:") for token in match.split()}
        if tokens & _NON_NAME_WORDS:
            continue
        matches.append(match)
    return matches


def _looks_like_person_roster(line: str) -> bool:
    """Detect personnel lists such as 'Name (Founder), Name (Eng. Head)' or bare name lists."""
    role_hits = len(_ROLE_PAREN_RE.findall(line))
    if role_hits >= 2:
        return True

    names = _personal_name_matches(line)
    # A single role tag alongside several personal names is still a roster.
    if role_hits >= 1 and len(names) >= 2:
        return True

    # A bare comma-separated list of people, e.g. "Priya Nair, Harish Kumar, Vishwas Kulkarni".
    if len(names) >= 3 and line.count(",") >= 2:
        covered = sum(len(name) for name in names)
        if covered >= len(line) * 0.5:
            return True

    return False


def _looks_like_noise_line(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return True
    if any(pattern.match(stripped) for pattern in _NOISE_LINE_PATTERNS):
        return True
    if _looks_like_person_roster(stripped):
        return True
    # Title-block leftovers embed a 'Topic:' marker mid-line; the real subject is kept as the heading.
    if re.search(r"\btopic\s*:", stripped, re.IGNORECASE):
        return True
    if re.match(r"^project\s+(owner|manager|coordinator|name)\b", stripped, re.IGNORECASE):
        return True
    lowered = stripped.lower()
    if _MEETING_CHATTER_RE.search(stripped):
        return True
    if re.search(r"\b[a-z]\s+days\b", lowered):
        return True
    if len(stripped) <= 2:
        return True
    if stripped.count("?") >= 2:
        return True
    if len(_tokenize(stripped)) == 0 and not re.search(r"\d", stripped):
        return True
    return False


def _clean_extracted_text(text: str) -> str:
    seen: set[str] = set()
    cleaned_lines: list[str] = []
    for raw_line in text.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if _looks_like_noise_line(line):
            continue
        normalized_key = line.lower()
        if normalized_key in seen:
            continue
        seen.add(normalized_key)
        cleaned_lines.append(line)
    return "\n".join(cleaned_lines)


def _safe_read_text(path: Path, max_chars: int) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")[:max_chars]


def _load_embedding_model():
    global _EMBEDDING_MODEL
    if _EMBEDDING_MODEL is not None:
        return _EMBEDDING_MODEL
    try:
        from sentence_transformers import SentenceTransformer
    except Exception:
        return None

    _EMBEDDING_MODEL = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")
    return _EMBEDDING_MODEL


def _read_docx(path: Path, max_chars: int) -> str:
    doc = Document(str(path))
    lines: list[str] = []
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        if paragraph.text:
            lines.append(f"[Paragraph {idx}] {paragraph.text}")
        if sum(len(line) for line in lines) >= max_chars:
            break
    text = _clean_extracted_text("\n".join(lines))
    return text[:max_chars]


def _read_pdf(path: Path, max_chars: int) -> str:
    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            chunks.append(f"[Page {page_number}]\n{page_text}")
        if sum(len(c) for c in chunks) >= max_chars:
            break
    return _clean_extracted_text("\n".join(chunks))[:max_chars]


def _iter_pptx_shapes(shapes) -> Iterable:
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_pptx_shapes(shape.shapes)


def _pptx_shape_text(shape, include_tables: bool = True) -> list[str]:
    lines: list[str] = []

    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = " ".join(run.text.strip() for run in paragraph.runs if run.text and run.text.strip()).strip()
            if text:
                lines.append(text)

    if include_tables and getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return lines


def _shape_table(shape, max_rows: int = 18, max_cols: int = 6) -> list[list[str]] | None:
    """Return a slide table as a trimmed row/cell matrix, or None when not usable."""
    if not getattr(shape, "has_table", False):
        return None

    try:
        table = shape.table
    except Exception:
        return None

    rows: list[list[str]] = []
    for row in list(table.rows)[:max_rows]:
        cells = [re.sub(r"\s+", " ", cell.text or "").strip() for cell in row.cells][:max_cols]
        if any(cells):
            rows.append(cells)

    if len(rows) < 2:
        return None

    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]

    # Drop columns that are entirely empty so narrow index columns do not distort layout.
    keep = [i for i in range(width) if any(row[i] for row in normalized)]
    if not keep:
        return None
    trimmed = [[row[i] for i in keep] for row in normalized]
    if len(trimmed[0]) < 2:
        return None
    return trimmed


def _normalize_slide_title(raw_title: str, max_length: int = 100) -> str:
    """Reduce noisy placeholder titles to the meaningful topic they describe."""
    title = re.sub(r"\s+", " ", raw_title or "").strip()
    if not title:
        return ""

    # Several decks append the real subject after a 'Topic:' marker inside a brainstorm title.
    if re.search(r"\btopic\s*:", title, re.IGNORECASE):
        candidate = re.split(r"\btopic\s*:", title, flags=re.IGNORECASE)[-1].strip()
        if len(candidate) >= 4:
            title = candidate

    title = title.strip(" -–—:|")
    if len(title) > max_length:
        clipped = title[:max_length]
        if " " in clipped:
            clipped = clipped.rsplit(" ", 1)[0]
        title = clipped.rstrip(" ,;:-–—") + "…"
    return title


def _read_pptx(path: Path, max_chars: int) -> str:
    prs = Presentation(str(path))
    chunks: list[str] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = [f"Slide {slide_number}"]
        title_shape = slide.shapes.title
        if title_shape is not None and getattr(title_shape, "text", "").strip():
            slide_lines.append(f"Title: {title_shape.text.strip()}")

        picture_count = 0
        for shape in _iter_pptx_shapes(slide.shapes):
            slide_lines.extend(_pptx_shape_text(shape))
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1

        if picture_count:
            slide_lines.append(f"Visual elements present: {picture_count}")

        if slide_lines:
            chunks.append("\n".join(slide_lines))
        if sum(len(c) for c in chunks) >= max_chars:
            break
    return _clean_extracted_text("\n".join(chunks))[:max_chars]


def _mime_type_from_ext(ext: str) -> str:
    normalized = ext.lower().lstrip(".")
    if normalized in {"jpg", "jpeg"}:
        return "image/jpeg"
    if normalized == "png":
        return "image/png"
    if normalized == "gif":
        return "image/gif"
    if normalized == "bmp":
        return "image/bmp"
    if normalized == "tif" or normalized == "tiff":
        return "image/tiff"
    if normalized == "webp":
        return "image/webp"
    return "application/octet-stream"


RASTER_IMAGE_EXTENSIONS = ("png", "jpg", "jpeg", "webp")
DEFAULT_MIN_IMAGE_BYTES = 12000


@dataclass
class SlideRecord:
    slide_number: int
    title: str
    lines: list[str]
    images: list[ImageInput]
    tables: list[list[list[str]]] = field(default_factory=list)


def _slide_images(
    source_stem: str,
    slide,
    slide_number: int,
    seen_hashes: set[str],
    min_image_bytes: int,
    max_images_per_slide: int = 3,
) -> list[ImageInput]:
    """Collect raster figures attached to a single slide, preserving slide order."""
    images: list[ImageInput] = []
    figure_index = 0

    for rel in slide.part.rels.values():
        if "image" not in getattr(rel, "reltype", ""):
            continue

        target_ref = getattr(rel, "target_ref", "") or ""
        ext = target_ref.rsplit(".", 1)[-1].lower() if "." in target_ref else "png"
        if ext not in RASTER_IMAGE_EXTENSIONS:
            continue

        try:
            blob = rel.target_part.blob
        except Exception:
            continue

        if len(blob) < min_image_bytes:
            continue

        digest = hashlib.sha1(blob).hexdigest()
        if digest in seen_hashes:
            continue
        seen_hashes.add(digest)

        figure_index += 1
        images.append(
            ImageInput(
                image_name=f"{source_stem} - Slide {slide_number} - Figure {figure_index}.{ext}",
                mime_type=_mime_type_from_ext(ext),
                content_base64=base64.b64encode(blob).decode("ascii"),
            )
        )
        if figure_index >= max_images_per_slide:
            break

    return images


def extract_pptx_slide_records(
    path: Path,
    max_slides: int = 400,
    max_total_images: int = 240,
    min_image_bytes: int = DEFAULT_MIN_IMAGE_BYTES,
) -> list[SlideRecord]:
    """Return every slide in presentation order with its cleaned text and figures."""
    try:
        prs = Presentation(str(path))
    except Exception:
        return []

    records: list[SlideRecord] = []
    seen_hashes: set[str] = set()
    total_images = 0

    for slide_number, slide in enumerate(prs.slides, start=1):
        if slide_number > max_slides:
            break

        title = ""
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None
        if title_shape is not None and getattr(title_shape, "text", "").strip():
            title = _normalize_slide_title(title_shape.text)

        raw_lines: list[str] = []
        tables: list[list[list[str]]] = []
        for shape in _iter_pptx_shapes(slide.shapes):
            raw_lines.extend(_pptx_shape_text(shape, include_tables=False))
            table = _shape_table(shape)
            if table is not None:
                tables.append(table)

        lines: list[str] = []
        seen_lines: set[str] = set()
        for raw_line in raw_lines:
            line = re.sub(r"\s+", " ", raw_line).strip()
            # Stray asterisks from source formatting would render as literal markdown markers.
            line = re.sub(r"\*{2,}", " ", line)
            line = re.sub(r"\s+", " ", line).strip()
            if not line or _looks_like_noise_line(line):
                continue
            key = line.lower()
            if key in seen_lines or key == title.lower():
                continue
            seen_lines.add(key)
            lines.append(line)

        images: list[ImageInput] = []
        if total_images < max_total_images:
            images = _slide_images(path.stem, slide, slide_number, seen_hashes, min_image_bytes)
            total_images += len(images)

        if not lines and not images and not title and not tables:
            continue

        records.append(
            SlideRecord(
                slide_number=slide_number,
                title=title,
                lines=lines,
                images=images,
                tables=tables,
            )
        )

    return records


def _extract_pptx_images(
    path: Path,
    max_images: int,
    seen_hashes: set[str],
    min_image_bytes: int = DEFAULT_MIN_IMAGE_BYTES,
) -> list[ImageInput]:
    records = extract_pptx_slide_records(
        path,
        max_total_images=max_images,
        min_image_bytes=min_image_bytes,
    )

    images: list[ImageInput] = []
    for record in records:
        for image in record.images:
            digest = hashlib.sha1(base64.b64decode(image.content_base64)).hexdigest()
            seen_hashes.add(digest)
            images.append(image)
            if len(images) >= max_images:
                return images
    return images


def _extract_pdf_images(path: Path, max_images: int, seen_hashes: set[str]) -> list[ImageInput]:
    reader = PdfReader(str(path))
    images: list[ImageInput] = []

    for page_number, page in enumerate(reader.pages, start=1):
        image_index = 0
        try:
            page_images = page.images
        except Exception:
            continue

        for img_name, img_obj in page_images.items():
            blob = getattr(img_obj, "data", None)
            if not blob or len(blob) < 8000:
                continue
            digest = hashlib.sha1(blob).hexdigest()
            if digest in seen_hashes:
                continue
            seen_hashes.add(digest)
            image_index += 1
            name = getattr(img_obj, "name", "") or img_name
            ext = name.split(".")[-1].lower() if "." in name else "png"
            if ext not in ("png", "jpg", "jpeg", "webp"):
                ext = "png"
            mime_type = _mime_type_from_ext(ext)
            images.append(
                ImageInput(
                    image_name=f"{path.stem} - Page {page_number} - Visual {image_index}.{ext}",
                    mime_type=mime_type,
                    content_base64=base64.b64encode(blob).decode("ascii"),
                )
            )
            if len(images) >= max_images:
                return images

    return images


def extract_workspace_images_from_paths(
    paths: list[Path],
    max_total_images: int = 12,
    min_image_bytes: int = DEFAULT_MIN_IMAGE_BYTES,
) -> list[ImageInput]:
    if max_total_images <= 0:
        return []

    images: list[ImageInput] = []
    seen_hashes: set[str] = set()
    for path in paths:
        remaining = max_total_images - len(images)
        if remaining <= 0:
            break
        suffix = path.suffix.lower()
        try:
            if suffix == ".pptx":
                images.extend(
                    _extract_pptx_images(
                        path,
                        max_images=remaining,
                        seen_hashes=seen_hashes,
                        min_image_bytes=min_image_bytes,
                    )
                )
            elif suffix == ".pdf":
                images.extend(_extract_pdf_images(path, max_images=remaining, seen_hashes=seen_hashes))
        except Exception:
            continue
    return images


def read_document_text(path: Path, max_chars: int = 6000) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return _safe_read_text(path, max_chars)
    if suffix == ".docx":
        return _read_docx(path, max_chars)
    if suffix == ".pdf":
        return _read_pdf(path, max_chars)
    if suffix == ".pptx":
        return _read_pptx(path, max_chars)
    return ""


def _read_document_text_cached(path: Path, max_chars: int = 6000) -> tuple[str, bool]:
    stat = path.stat()
    cache_key = (str(path.resolve()), max_chars)
    fingerprint = (stat.st_mtime_ns, stat.st_size, max_chars)
    cached = _WORKSPACE_TEXT_CACHE.get(cache_key)
    if cached and cached[0] == fingerprint:
        return cached[1], True

    text = read_document_text(path, max_chars=max_chars)
    _WORKSPACE_TEXT_CACHE[cache_key] = (fingerprint, text)
    return text, False


def _iter_workspace_files(root: Path) -> Iterable[Path]:
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            if any(part.startswith(".") for part in path.parts):
                continue
            yield path


def _score_file(path: Path, prompt_tokens: set[str], hints: list[str]) -> int:
    name_tokens = _tokenize(path.name)
    score = len(prompt_tokens.intersection(name_tokens))

    lowered = str(path).lower()
    for hint in hints:
        if hint and hint.lower() in lowered:
            score += 5

    if path.suffix.lower() in {".pptx", ".pdf"}:
        score += 1

    return score


def select_workspace_files(workspace_root: Path, user_prompt: str, hints: list[str], max_docs: int = 5) -> list[Path]:
    prompt_tokens = _tokenize(user_prompt)
    candidates = list(_iter_workspace_files(workspace_root))
    ranked = sorted(candidates, key=lambda p: _score_file(p, prompt_tokens, hints), reverse=True)

    selected: list[Path] = []
    for path in ranked:
        if len(selected) >= max_docs:
            break
        if _score_file(path, prompt_tokens, hints) <= 0 and hints:
            continue
        selected.append(path)

    if not selected and candidates:
        selected = ranked[: min(max_docs, len(ranked))]
    return selected


def _chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    normalized = "\n".join(line.strip() for line in text.splitlines())
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", normalized) if p.strip()]
    if not paragraphs:
        paragraphs = [normalized.strip()] if normalized.strip() else []

    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        candidate = f"{current}\n\n{para}".strip() if current else para
        if len(candidate) <= chunk_size:
            current = candidate
            continue

        if current:
            chunks.append(current)

        if len(para) <= chunk_size:
            current = para
            continue

        start = 0
        while start < len(para):
            end = min(start + chunk_size, len(para))
            chunks.append(para[start:end])
            if end == len(para):
                break
            start = max(0, end - overlap)
        current = ""

    if current:
        chunks.append(current)

    return chunks


def _idf(chunks: list[str], token: str) -> float:
    df = sum(1 for chunk in chunks if token in _tokenize(chunk))
    return math.log((1 + len(chunks)) / (1 + df)) + 1.0


def _lexical_score(query_tokens: set[str], chunk: str, all_chunks: list[str]) -> float:
    if not query_tokens:
        return 0.0
    chunk_tokens = _tokenize(chunk)
    if not chunk_tokens:
        return 0.0
    score = 0.0
    for token in query_tokens:
        if token in chunk_tokens:
            score += _idf(all_chunks, token)
    return score / max(len(chunk_tokens), 1)


def _embedding_scores(query: str, chunks: list[str]) -> list[float]:
    model = _load_embedding_model()
    if model is None:
        return [0.0 for _ in chunks]
    try:
        import numpy as np

        q = model.encode([query], normalize_embeddings=True)[0]
        c = model.encode(chunks, normalize_embeddings=True)
        sim = np.dot(c, q)
        return [float(x) for x in sim]
    except Exception:
        return [0.0 for _ in chunks]


def _pairwise_token_overlap(left: str, right: str) -> float:
    left_tokens = _tokenize(left)
    right_tokens = _tokenize(right)
    if not left_tokens or not right_tokens:
        return 0.0
    intersection = len(left_tokens.intersection(right_tokens))
    union = len(left_tokens.union(right_tokens))
    return intersection / union if union else 0.0


def _select_diverse_top_chunks(items: list[RetrievedChunk], top_k: int, diversity_lambda: float = 0.82) -> list[RetrievedChunk]:
    if len(items) <= top_k:
        return sorted(items, key=lambda x: x.score, reverse=True)

    remaining = sorted(items, key=lambda x: x.score, reverse=True)
    selected: list[RetrievedChunk] = []

    while remaining and len(selected) < top_k:
        if not selected:
            selected.append(remaining.pop(0))
            continue

        best_index = 0
        best_value = float("-inf")
        for idx, candidate in enumerate(remaining):
            redundancy = max(_pairwise_token_overlap(candidate.text, chosen.text) for chosen in selected)
            mmr_value = (diversity_lambda * candidate.score) - ((1.0 - diversity_lambda) * redundancy)
            if mmr_value > best_value:
                best_value = mmr_value
                best_index = idx
        selected.append(remaining.pop(best_index))

    return sorted(selected, key=lambda x: x.score, reverse=True)


def build_workspace_context(
    workspace_root: Path,
    user_prompt: str,
    hints: list[str],
    max_docs: int = 5,
    max_chars_total: int = 15000,
    max_chars_per_doc: int = 12000,
    chunk_size_chars: int = 900,
    chunk_overlap_chars: int = 180,
    top_k_chunks: int = 8,
    use_embeddings: bool = True,
) -> tuple[str, list[str], list[RetrievedChunk], RetrievalStats]:
    start = time.perf_counter()
    candidates = list(_iter_workspace_files(workspace_root))
    selected = select_workspace_files(workspace_root, user_prompt, hints, max_docs=max_docs)

    raw_chunks: list[tuple[str, str]] = []
    cache_hits = 0
    cache_misses = 0
    for path in selected:
        text, was_cached = _read_document_text_cached(path, max_chars=max_chars_per_doc)
        if was_cached:
            cache_hits += 1
        else:
            cache_misses += 1
        if not text.strip():
            continue
        split = _chunk_text(text.strip(), chunk_size=chunk_size_chars, overlap=chunk_overlap_chars)
        for idx, chunk in enumerate(split):
            raw_chunks.append((f"{path.name}#c{idx+1}", chunk))

    if not raw_chunks:
        return "", [], [], RetrievalStats(
            candidate_files=len(candidates),
            selected_files=len(selected),
            chunks_scored=0,
            returned_chunks=0,
            cache_hits=cache_hits,
            cache_misses=cache_misses,
            context_chars=0,
            retrieval_latency_ms=int((time.perf_counter() - start) * 1000),
            top_score=0.0,
            avg_top_score=0.0,
            used_embeddings=use_embeddings,
        )

    chunk_texts = [t for _, t in raw_chunks]
    prompt_tokens = _tokenize(user_prompt)
    lexical = [_lexical_score(prompt_tokens, chunk, chunk_texts) for chunk in chunk_texts]
    embedding = _embedding_scores(user_prompt, chunk_texts) if use_embeddings else [0.0 for _ in chunk_texts]

    hints_lower = [h.lower() for h in hints if h.strip()]
    scored: list[RetrievedChunk] = []
    for i, (chunk_id, chunk_text) in enumerate(raw_chunks):
        source_name = chunk_id.split("#", 1)[0]
        hint_bonus = 0.0
        for hint in hints_lower:
            if hint in source_name.lower():
                hint_bonus += 0.15
        hybrid = (0.65 * embedding[i]) + (0.35 * lexical[i]) + hint_bonus
        scored.append(
            RetrievedChunk(
                source=source_name,
                chunk_id=chunk_id,
                score=round(hybrid, 6),
                text=chunk_text,
            )
        )

    top = _select_diverse_top_chunks(scored, top_k_chunks)

    snippets: list[str] = []
    sources: list[str] = []
    consumed = 0
    for item in top:
        source_line = f"Source Chunk: {item.chunk_id} (score={item.score})"
        chunk = f"{source_line}\n{item.text.strip()}"
        if consumed + len(chunk) > max_chars_total:
            remaining = max_chars_total - consumed
            if remaining <= 0:
                break
            chunk = chunk[:remaining]
        snippets.append(chunk)
        if item.source not in sources:
            sources.append(item.source)
        consumed += len(chunk)
        if consumed >= max_chars_total:
            break

    top_scores = [item.score for item in top]
    context = "\n\n".join(snippets)
    stats = RetrievalStats(
        candidate_files=len(candidates),
        selected_files=len(selected),
        chunks_scored=len(scored),
        returned_chunks=len(top),
        cache_hits=cache_hits,
        cache_misses=cache_misses,
        context_chars=len(context),
        retrieval_latency_ms=int((time.perf_counter() - start) * 1000),
        top_score=round(max(top_scores), 6) if top_scores else 0.0,
        avg_top_score=round(sum(top_scores) / len(top_scores), 6) if top_scores else 0.0,
        used_embeddings=use_embeddings,
    )
    return context, sources, top, stats

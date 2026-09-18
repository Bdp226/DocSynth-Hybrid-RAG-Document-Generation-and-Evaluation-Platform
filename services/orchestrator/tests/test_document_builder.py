from io import BytesIO

from docx import Document as DocxDocument

from app.document_builder import generate_artifacts
from app.models import ImageInput

ONE_PIXEL_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def test_generate_artifacts_builds_pdf_and_docx() -> None:
    artifacts = generate_artifacts("doc-1", "Hello\n\nWorld", ["pdf", "docx"])

    formats = {a.format for a in artifacts}
    assert formats == {"pdf", "docx"}
    assert all(len(a.content) > 0 for a in artifacts)


def test_docx_strips_markdown_bold_markers_and_adds_brand_header() -> None:
    text = "**Introduction**\n\nThis has **bold** words.\n\n- Point one\n- **Point two**"
    artifacts = generate_artifacts("doc-2", text, ["docx"])
    docx_artifact = next(a for a in artifacts if a.format == "docx")

    doc = DocxDocument(BytesIO(docx_artifact.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "**" not in full_text
    assert "Introduction" in full_text
    assert "bold" in full_text

    header_text = doc.sections[0].header.paragraphs[0].text
    assert "SIEMENS" in header_text
    assert "Healthineers" in header_text


def test_inline_image_tokens_render_in_body_and_skip_appendix() -> None:
    images = [
        ImageInput(
            image_name="deck - Slide 1 - Figure 1.png",
            mime_type="image/png",
            content_base64=ONE_PIXEL_PNG_B64,
        ),
        ImageInput(
            image_name="unused-figure.png",
            mime_type="image/png",
            content_base64=ONE_PIXEL_PNG_B64,
        ),
    ]
    text = (
        "# Deck Documentation\n\n"
        "## Slide 1: Overview\n\n"
        "[[IMAGE:deck - Slide 1 - Figure 1.png]]\n\n"
        "This section explains the slide in formal prose."
    )

    artifacts = generate_artifacts("doc-3", text, ["docx", "pdf"], image_inputs=images)
    docx_artifact = next(a for a in artifacts if a.format == "docx")
    pdf_artifact = next(a for a in artifacts if a.format == "pdf")

    doc = DocxDocument(BytesIO(docx_artifact.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert "[[IMAGE:" not in full_text
    # Captions are sequential and never expose source slide numbering.
    assert "Figure 1 — Slide 1: Overview" in full_text
    assert "Slide 1 - Figure 1.png" not in full_text
    # The inline image is not repeated in the appendix, but the unused one still is.
    assert "Image Appendix" in full_text
    assert "unused-figure.png" in full_text

    assert len(doc.inline_shapes) == 2
    assert len(pdf_artifact.content) > 0


def test_markdown_tables_render_as_real_tables() -> None:
    text = (
        "# Costing\n\n"
        "## Budget\n\n"
        "| Item | Cost (USD) |\n"
        "| --- | --- |\n"
        "| Setup | 1200 |\n"
        "| Support | 800 |\n\n"
        "Closing prose."
    )

    artifacts = generate_artifacts("doc-4", text, ["docx", "pdf"])
    docx_artifact = next(a for a in artifacts if a.format == "docx")
    pdf_artifact = next(a for a in artifacts if a.format == "pdf")

    doc = DocxDocument(BytesIO(docx_artifact.content))
    assert len(doc.tables) == 1

    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Item"
    assert table.cell(0, 1).text == "Cost (USD)"
    assert table.cell(1, 0).text == "Setup"
    assert table.cell(2, 1).text == "800"

    # The divider row must not survive as content.
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "---" not in body_text
    assert "| Item |" not in body_text
    assert len(pdf_artifact.content) > 0


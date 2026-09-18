from __future__ import annotations

import base64
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.workspace_context import build_workspace_context
from app.workspace_context import extract_pptx_slide_records
from app.workspace_context import extract_workspace_images_from_paths
from app.workspace_context import _looks_like_noise_line, _normalize_slide_title


def test_person_rosters_are_filtered_but_domain_phrases_are_kept() -> None:
    rosters = [
        "Aditi Kothari (Founder), Jagadeesh Dyaberi (Eng. Head)",
        "Yash Krishan (PL & SM), Deepesh Genani, Bharath Kumar K, Nihit Gupta",
        "Paranthaman Ethiraj, Harish Kumar B K, Vishwas Kulkarni, Kishan K S",
    ]
    for line in rosters:
        assert _looks_like_noise_line(line), line

    keepers = [
        "Solution Architecture, Data Flow, Environment Setup",
        "Sandbox Architecture, Network Design, Security Controls",
        "Requesting SSL Certificate for the sandbox domain",
    ]
    for line in keepers:
        assert not _looks_like_noise_line(line), line


def test_documentation_prose_is_not_treated_as_meeting_chatter() -> None:
    # Bare reporting verbs are legitimate in documentation prose.
    assert not _looks_like_noise_line("The figures are explained in the following section.")
    assert not _looks_like_noise_line("No external assumptions have been introduced.")
    # A named participant performing the verb is real chatter.
    assert _looks_like_noise_line("Kaushik explained the sandbox requirements to the team.")
    assert _looks_like_noise_line("Looping in SCM team for the pending approval.")


def test_normalize_slide_title_prefers_topic_and_caps_length() -> None:
    raw = (
        'Aiden (AI + "aide"), or Continuum : Evokes an unbroken thread of context and execution '
        "— sounds premium in an exec deck Topic: Intelli Hub – AI based dashboards"
    )
    assert _normalize_slide_title(raw) == "Intelli Hub – AI based dashboards"

    long_title = "A " + ("very " * 60) + "long title"
    normalized = _normalize_slide_title(long_title)
    assert len(normalized) <= 101
    assert normalized.endswith("…")


def test_build_workspace_context_uses_hint_and_source(tmp_path: Path) -> None:
    (tmp_path / "siemens_style.md").write_text("Tone: clinical concise style.", encoding="utf-8")
    (tmp_path / "other.txt").write_text("random content", encoding="utf-8")

    context, sources, retrieved, stats = build_workspace_context(
        workspace_root=tmp_path,
        user_prompt="create a clinical summary",
        hints=["siemens_style"],
        max_docs=2,
        max_chars_total=2000,
        use_embeddings=False,
    )

    assert "clinical concise style" in context
    assert "siemens_style.md" in sources
    assert retrieved
    assert stats.selected_files >= 1
    assert stats.returned_chunks >= 1


def test_build_workspace_context_reports_cache_hits_on_repeated_reads(tmp_path: Path) -> None:
    doc_path = tmp_path / "siemens_style.md"
    doc_path.write_text("clinical concise content for repeated retrieval", encoding="utf-8")

    first = build_workspace_context(
        workspace_root=tmp_path,
        user_prompt="clinical concise summary",
        hints=["siemens_style"],
        use_embeddings=False,
    )
    second = build_workspace_context(
        workspace_root=tmp_path,
        user_prompt="clinical concise summary",
        hints=["siemens_style"],
        use_embeddings=False,
    )

    assert first[3].cache_misses >= 1
    assert second[3].cache_hits >= 1


def test_build_workspace_context_diversifies_redundant_chunks(tmp_path: Path) -> None:
    (tmp_path / "primary.md").write_text(
        "Alpha rollout summary.\n\n" + ("same repeated cluster text " * 50) + "\n\nUnique governance section with approval workflow.",
        encoding="utf-8",
    )

    context, _, retrieved, stats = build_workspace_context(
        workspace_root=tmp_path,
        user_prompt="alpha rollout governance approval workflow",
        hints=["primary"],
        chunk_size_chars=180,
        chunk_overlap_chars=30,
        top_k_chunks=2,
        use_embeddings=False,
    )

    assert context
    assert len(retrieved) == 2
    assert stats.returned_chunks == 2


def test_extract_workspace_images_from_pptx(tmp_path: Path) -> None:
    # A realistic slide graphic: content triage deliberately discards 1x1 pixel
    # spacers, icons and other non-informative assets, so the fixture has to be a
    # diagram the pipeline would actually publish.
    canvas = Image.new("RGB", (600, 400), (255, 255, 255))
    draw = ImageDraw.Draw(canvas)
    for x in range(0, 600, 50):
        draw.line([(x, 0), (x, 400)], fill=(20, 40, 90), width=2)
    for y in range(0, 400, 50):
        draw.line([(0, y), (600, y)], fill=(200, 60, 30), width=2)
    for x in range(0, 600, 100):
        draw.rectangle([x + 2, 2, x + 48, 398], outline=(0, 0, 0), width=1)

    image_path = tmp_path / "slide.png"
    canvas.save(image_path, format="PNG")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(0.6)).text_frame.text = "Sandbox architecture"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1.2), width=Inches(3), height=Inches(2))
    deck_path = tmp_path / "deck.pptx"
    prs.save(str(deck_path))

    images = extract_workspace_images_from_paths([deck_path], max_total_images=4, min_image_bytes=0)

    assert len(images) == 1
    assert images[0].image_name.startswith("deck - Slide 1 - Figure 1")
    assert images[0].mime_type.startswith("image/")
    assert images[0].content_base64

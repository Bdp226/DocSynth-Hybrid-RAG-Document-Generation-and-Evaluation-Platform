"""Introspect the real deck so restructuring decisions are grounded in actual content."""
from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

DECK = Path(__file__).resolve().parents[3] / "Sandbox environment.pptx"


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def main() -> None:
    prs = Presentation(str(DECK))
    titles = []
    table_slides = []
    roster_examples = []
    long_titles = 0

    for number, slide in enumerate(prs.slides, start=1):
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None
        title = ""
        if title_shape is not None and getattr(title_shape, "text", "").strip():
            title = re.sub(r"\s+", " ", title_shape.text).strip()
        titles.append((number, title))
        if len(title) > 90:
            long_titles += 1

        for shape in iter_shapes(slide.shapes):
            if getattr(shape, "has_table", False):
                tbl = shape.table
                rows = len(tbl.rows)
                cols = len(tbl.columns)
                header = [c.text.strip() for c in tbl.rows[0].cells][:6]
                table_slides.append((number, rows, cols, header))

            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    text = " ".join(r.text.strip() for r in para.runs if r.text.strip()).strip()
                    if not text:
                        continue
                    if len(re.findall(r"\([A-Za-z][A-Za-z .&]{0,14}\)", text)) >= 2:
                        roster_examples.append((number, text[:150]))

    print(f"TOTAL_SLIDES={len(titles)}")
    print(f"LONG_TITLES(>90 chars)={long_titles}")

    counts = Counter(t for _, t in titles if t)
    print("\n--- MOST REPEATED TITLES ---")
    for title, count in counts.most_common(8):
        print(f"{count:3d}x  {title[:110]}")

    print("\n--- FIRST 25 TITLES ---")
    for number, title in titles[:25]:
        print(f"{number:3d}: {title[:110]}")

    print(f"\n--- TABLES FOUND: {len(table_slides)} ---")
    for number, rows, cols, header in table_slides[:15]:
        print(f"slide {number:3d}  {rows}x{cols}  header={header}")

    print(f"\n--- ROSTER-LIKE LINES: {len(roster_examples)} ---")
    for number, text in roster_examples[:12]:
        print(f"slide {number:3d}: {text}")


if __name__ == "__main__":
    main()
